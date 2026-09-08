"""Write an ATS-friendly DOCX from the tailor markdown convention.

Understands: # name, a contact line right under it, ## section headings,
- / * bullets, --- rules, a whole-line _italic_ disclaimer, and inline
**bold** anywhere (single column, no tables — ATS parsers choke on those).
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor

_INK = RGBColor(0x11, 0x18, 0x27)
_MUTED = RGBColor(0x5B, 0x66, 0x6E)
_RULE = "C8CDD2"
_BOLD = re.compile(r"\*\*(.+?)\*\*")


def write_markdown_docx(text: str, dest: str | Path) -> str:
    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    lines = (text or "").splitlines()
    i, n = 0, len(lines)
    while i < n:
        stripped = lines[i].strip()
        if not stripped:
            i += 1
            continue

        if stripped.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _emit(p, stripped[2:].strip(), size=Pt(20), bold=True, color=_INK)
            i += 1
            j = i
            while j < n and not lines[j].strip():
                j += 1
            if j < n and not lines[j].strip().startswith(("#", "- ", "* ", "---")):
                cp = doc.add_paragraph()
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cp.paragraph_format.space_after = Pt(10)
                _emit(cp, lines[j].strip(), size=Pt(10), color=_MUTED)
                _bottom_rule(cp)
                i = j + 1
            continue

        if stripped == "---":
            hp = doc.add_paragraph()
            hp.paragraph_format.space_after = Pt(4)
            _bottom_rule(hp)
            i += 1
            continue

        if stripped.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            _emit(p, stripped[3:].strip().upper(), size=Pt(12), bold=True, color=_INK)
            _bottom_rule(p, size=4)
            i += 1
            continue

        if stripped.startswith(("- ", "* ")):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            _emit(p, stripped[2:].strip())
            i += 1
            continue

        if len(stripped) > 2 and stripped.startswith("_") and stripped.endswith("_"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            _emit(p, stripped[1:-1].strip(), size=Pt(9), color=_MUTED, italic=True)
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        if stripped.startswith("**"):
            p.paragraph_format.space_before = Pt(8)
        _emit(p, stripped)
        i += 1

    doc.save(dest)
    return str(dest.resolve())


def _emit(paragraph, text: str, *, size=Pt(11), bold=False, color=None, italic=False) -> None:
    pos = 0
    for m in _BOLD.finditer(text):
        if m.start() > pos:
            _run(paragraph, text[pos:m.start()], size, bold, color, italic)
        _run(paragraph, m.group(1), size, True, color, italic)
        pos = m.end()
    if pos < len(text):
        _run(paragraph, text[pos:], size, bold, color, italic)


def _run(paragraph, text: str, size, bold: bool, color, italic: bool) -> None:
    if not text:
        return
    run = paragraph.add_run(text)
    run.font.name = "Calibri"
    run.font.size = size
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _bottom_rule(paragraph, size: int = 6) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), _RULE)
    borders.append(bottom)
    p_pr.append(borders)
