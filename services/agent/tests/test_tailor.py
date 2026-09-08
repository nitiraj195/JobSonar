from pathlib import Path

from jobsonar_agent.llm.factory import resolve_tailor_llm
from jobsonar_agent.llm.fake import FakeLLM
from jobsonar_agent.tailor.docx import write_markdown_docx
from jobsonar_agent.tailor.fallback import fallback_draft
from jobsonar_agent.tailor.prompt import build_tailor_prompt, parse_tailor


def test_fake_llm_tailor_is_parseable():
    prompt = build_tailor_prompt(
        profile={"skills": ["kubernetes"], "seniority": "senior", "location": "Pune"},
        resume_text="DevOps Engineer. Kubernetes, Terraform.",
        jd={"title": "SRE", "company": "Acme", "jd_md": "Need kubernetes on-call."},
    )
    raw = FakeLLM().complete(prompt)
    parsed = parse_tailor(raw)
    assert parsed is not None
    resume, cover, notes = parsed
    assert "draft" in resume.lower()
    assert "hiring manager" in cover.lower()
    assert "do not submit" in notes.lower()


def test_parse_tailor_fence():
    raw = """```json
{"resume_md": "R", "cover_letter_md": "C", "notes_md": "N"}
```"""
    assert parse_tailor(raw) == ("R", "C", "N")


def test_parse_tailor_empty():
    assert parse_tailor("") is None
    assert parse_tailor("not useful") is None


def test_tailor_prompt_caps_and_forbids_apply():
    prompt = build_tailor_prompt(
        profile={"skills": ["aws"]},
        resume_text="x" * 20000,
        jd={"title": "SRE", "company": "Acme", "jd_md": "y" * 20000},
    )
    assert "Do not apply" in prompt
    assert "Do not invent" in prompt
    assert "x" * 12000 in prompt
    assert "x" * 12001 not in prompt
    assert "y" * 8000 in prompt
    assert "y" * 8001 not in prompt


def test_fallback_keeps_original_resume_and_jd_keywords():
    resume, cover, notes = fallback_draft(
        resume_text="NITIRAJ PATNE\nnitiraj195@gmail.com\nDevOps at Dkatalis. Kubernetes, Terraform, AWS.",
        profile={"skills": ["kubernetes", "terraform", "aws"]},
        jd={"title": "SRE", "company": "Acme", "jd_md": "Need kubernetes and terraform on-call."},
    )
    assert "NITIRAJ" in resume
    assert "Dkatalis" in resume
    assert "kubernetes" in resume.lower()
    assert "Acme" in cover
    assert "invented" in notes.lower() or "not invented" in notes.lower()


def test_write_markdown_docx(tmp_path: Path):
    dest = tmp_path / "r.docx"
    path = write_markdown_docx("# Name\n## Skills\n- kubernetes\nBody line", dest)
    assert Path(path).is_file()
    assert Path(path).stat().st_size > 1000


def test_write_markdown_docx_renders_inline_bold_not_asterisks(tmp_path: Path):
    from docx import Document as ReadDocument

    resume_md = (
        "# Jane Doe\n"
        "jane@example.com · Remote\n\n"
        "## Experience\n"
        "**Acme Corp** — Engineer (2020–2022)\n"
        "- Shipped **kubernetes** rollouts.\n"
    )
    dest = tmp_path / "r.docx"
    path = write_markdown_docx(resume_md, dest)
    doc = ReadDocument(path)
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "**" not in all_text
    bold_runs = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    assert "Acme Corp" in bold_runs
    assert "kubernetes" in bold_runs


def test_fallback_segments_experience_and_education_sections():
    resume_text = (
        "Nitiraj Patne\n"
        "nitiraj195@gmail.com\n"
        "EXPERIENCE\n"
        "DevOps Engineer, Dkatalis 2021 - Present\n"
        "Ran kubernetes clusters in production.\n"
        "Automated terraform pipelines.\n"
        "EDUCATION\n"
        "B.Tech Computer Science, Pune University 2016 - 2020\n"
    )
    resume, _cover, _notes = fallback_draft(
        resume_text=resume_text,
        profile={"skills": ["kubernetes", "terraform"]},
        jd={"title": "SRE", "company": "Acme", "jd_md": "Need kubernetes and terraform."},
    )
    assert "## Experience" in resume
    assert "**DevOps Engineer, Dkatalis 2021 - Present**" in resume
    assert "- Ran kubernetes clusters in production." in resume
    assert "## Education" in resume
    assert "B.Tech Computer Science, Pune University 2016 - 2020" in resume


def test_resolve_tailor_llm_never_bedrock(monkeypatch):
    from jobsonar_agent import config

    monkeypatch.setattr(config, "DEEP_DIVE_BACKEND", "bedrock")
    monkeypatch.setattr(config, "DEEP_DIVE_OPT_IN", True)
    llm = resolve_tailor_llm()
    assert isinstance(llm, FakeLLM)
